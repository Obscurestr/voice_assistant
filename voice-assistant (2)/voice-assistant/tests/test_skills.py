from voice_assistant.skills import REGISTRY


def run(intent, text, ctx):
    return REGISTRY[intent](text, ctx)


def test_alarm_and_calendar(ctx):
    assert run("alarm", "wake me up at 6:45 am", ctx).text == "Alarm set for 6:45 AM."
    assert run("alarm", "set an alarm", ctx).text.startswith("What time")
    assert "06:45" in run("calendar", "what's on my calendar", ctx).text


def test_reminder_strips_time(ctx):
    resp = run("reminder", "remind me to call mom at 5 pm", ctx)
    assert resp.text == "Okay, I'll remind you to call mom at 5:00 PM."
    assert ctx.memory["reminders"] == [{"task": "call mom", "time": "17:00"}]


def test_calculator_and_conversion(ctx):
    assert run("calculator", "what is 7 times 6", ctx).text == "The answer is 42."
    assert "zero" in run("calculator", "9 divided by 0", ctx).text
    assert run("unit_conversion", "convert 10 kilometers to miles", ctx).text == (
        "10 kilometers is 6.214 miles."
    )


def test_open_website_bug_is_fixed(ctx, env):
    # The notebook opened YouTube for every request because of `or "you tube"`.
    assert run("open_website", "open github", ctx).text == "Opening Github."
    assert env.opened == ["https://github.com"]
    assert run("open_website", "open something", ctx).text.startswith("Which website")


def test_directions_without_destination_asks(ctx):
    # The notebook crashed here with NameError: 'location' is not defined.
    assert run("directions", "give me directions", ctx).text == "Where would you like to go?"
    assert "Airport" in run("directions", "directions to the airport", ctx).text


def test_weather_uses_api_then_falls_back(ctx, env):
    env.weather_payloads = [
        {"results": [{"name": "Belgrade", "latitude": 44.8, "longitude": 20.5}]},
        {"current": {"temperature_2m": 17.6, "weather_code": 2}},
    ]
    assert run("weather", "weather in belgrade", ctx).text == (
        "It's 18 degrees Celsius with partly cloudy in Belgrade."
    )
    offline = run("weather", "weather in paris", ctx)
    assert "couldn't reach" in offline.text and offline.url


def test_smart_home_state(ctx):
    assert run("smart_home", "turn on the lights in the kitchen", ctx).text == (
        "Turned on the lights in the kitchen."
    )
    run("smart_home", "turn off the tv", ctx)
    status = run("device_status", "what's on", ctx).text
    assert "Lights in kitchen is on" in status and "tv in home is off" in status


def test_lists(ctx):
    run("todo_task", "add finish report to my to do list", ctx)
    assert run("todo_task", "show my to do list", ctx).text == (
        "Your to-do list has: finish report."
    )
    assert run("shopping_list", "clear my shopping list", ctx).text.endswith("empty.")


def test_routine_greets_by_name(ctx):
    ctx.memory["name"] = "Sara"
    assert run("routine", "good morning", ctx).text.startswith("Good morning, Sara. It's 9:05 AM.")


def test_shopping_list_exports_to_word(ctx, env, tmp_path):
    from docx import Document

    opened = []
    ctx.documents_dir = tmp_path / "docs"
    ctx.open_file = opened.append
    ctx.memory["name"] = "Sara"
    assert "empty" in run("shopping_list", "save my shopping list as a word document", ctx).text

    resp = run("shopping_list", "Make a shopping list document with milk, eggs and bread", ctx)
    assert resp.text == "I saved your shopping list with 3 items as a Word document and opened it."
    path = tmp_path / "docs" / "shopping_list_2026-10-21.docx"
    assert opened == [path] and path.exists()

    doc = Document(path)
    assert doc.paragraphs[0].text == "Sara's Shopping List"
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows[1:] == [["☐", "1", "Milk"], ["☐", "2", "Eggs"], ["☐", "3", "Bread"]]


def test_adding_several_items_at_once(ctx):
    assert run("shopping_list", "add milk, eggs and bread to my shopping list", ctx).text == (
        "Added milk, eggs and bread to your shopping list."
    )
    assert ctx.memory["shopping_list"] == ["milk", "eggs", "bread"]
