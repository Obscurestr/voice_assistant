"""Word document export (python-docx)."""

from __future__ import annotations

import datetime as dt
import os
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

CHECKBOX = "☐"  # ☐


def write_shopping_list(
    items: list[str], folder: Path, now: dt.datetime, owner: str | None = None
) -> Path:
    """Write a printable checklist to <folder>/shopping_list_<date>.docx and return its path.

    Exporting twice on the same day overwrites that day's file, so the document always
    matches the current list.
    """
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"shopping_list_{now:%Y-%m-%d}.docx"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    title = f"{owner}'s Shopping List" if owner else "Shopping List"
    doc.add_heading(title, level=0)
    subtitle = doc.add_paragraph(
        f"{now:%A, %B} {now.day}, {now.year}  ·  {len(items)} item{'s' if len(items) != 1 else ''}"
    )
    subtitle.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ("", "#", "Item"), strict=True):
        cell.text = label
    for number, item in enumerate(items, start=1):
        row = table.add_row().cells
        row[0].text = CHECKBOX
        row[1].text = str(number)
        row[2].text = item[:1].upper() + item[1:]
    # Fixed widths: Word only honours them with autofit off and set on every cell.
    table.autofit = False
    widths = (Cm(1.2), Cm(1.2), Cm(13.4))
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width
    for run in (r for row in table.rows[1:] for r in row.cells[0].paragraphs[0].runs):
        run.font.size = Pt(16)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(f"Created by Voice Assistant at {now:%H:%M}").italic = True
    note.runs[0].font.size = Pt(9)

    doc.save(path)
    return path


def open_file(path: Path) -> None:
    """Open a file with the default application (Word on Windows)."""
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)  # type: ignore[attr-defined]  # noqa: S606
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])  # noqa: S603, S607
        else:
            subprocess.Popen(["xdg-open", str(path)])  # noqa: S603, S607
    except OSError:
        pass  # no default app (e.g. a headless server); the file is still saved
